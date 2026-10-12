"""Make the collaborator copy without rewriting the supplied manuscript.

Run with a Python environment providing python-docx. Citations are numbered in
first-use order and resolved from the adjacent supplied BibTeX bibliography.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
source = (ROOT / "paper.md").read_text()
metadata, body = source.split("---", 2)[1:]
bib = (ROOT / "paper.bib").read_text()
entries = {}
for entry in re.split(r"(?m)^@", bib)[1:]:
    key = re.match(r"\w+\{([^,]+),", entry).group(1)
    fields = {}
    for match in re.finditer(r"(?m)^\s*(\w+)\s*=\s*\{(.*?)(?:\},?\s*\n|\}\s*$)", entry, re.S):
        fields[match.group(1)] = match.group(2)
    entries[key] = fields
keys = list(dict.fromkeys(re.findall(r"@([A-Za-z0-9]+)", body)))
assert set(keys) <= entries.keys(), "Unresolved citation"
numbers = {key: i + 1 for i, key in enumerate(keys)}

def plain(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\[(@[^\]]+)\]", lambda m: "[" + ", ".join(str(numbers[k]) for k in re.findall(r"@([A-Za-z0-9]+)", m[1])) + "]", text)
    return text.replace("`", "").replace("**", "")

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 22), ('Heading 1', 14)]:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
title = re.search(r"(?m)^title: '(.+)'", metadata)[1]
doc.add_paragraph(title, 'Title')
doc.add_paragraph('Ty Tuff\nORCID 0000-0001-5249-5197\nEnvironmental Data Science Innovation & Impact Lab, Cooperative Institute for Research in Environmental Sciences, University of Colorado Boulder, Boulder, Colorado, USA\n12 October 2026')
code = False
for block in re.split(r"\n\s*\n", body.strip()):
    if block.startswith('## '):
        doc.add_paragraph(block[3:].strip(), 'Heading 1')
    elif block.startswith('!['):
        match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", block)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(str(ROOT / match[2]), width=Inches(6.5))
        drawing = run._r.xpath('.//wp:docPr')[0]
        drawing.set('descr', match[1])
        caption = doc.add_paragraph('Figure 1. ' + match[1], 'Caption')
    elif block.startswith('```'):
        p = doc.add_paragraph()
        r = p.add_run('\n'.join(block.splitlines()[1:-1]))
        r.font.name = 'Consolas'
        r.font.size = Pt(10)
    else:
        doc.add_paragraph(plain(block.replace('\n', ' ')))
doc.add_paragraph('References', 'Heading 1')
for key in keys:
    e = entries[key]
    author = e['author'].split(' and ')
    author = author[0] + (' et al.' if len(author) > 2 else (' and ' + author[1] if len(author) == 2 else ''))
    def clean(t):
        return t.replace('{', '').replace('}', '').replace('---', '—').replace('--', '–')
    reference = f"[{numbers[key]}] {clean(author)} ({e['year']}). {clean(e['title'])}."
    if 'journal' in e:
        reference += ' ' + clean(e['journal']) + '.'
    if 'booktitle' in e:
        reference += ' ' + clean(e['booktitle']) + '.'
    if 'doi' in e:
        reference += ' https://doi.org/' + e['doi']
    elif 'url' in e:
        reference += ' ' + e['url']
    if 'note' in e:
        reference += ' ' + clean(e['note'])
    doc.add_paragraph(reference)
doc.core_properties.title = title
doc.core_properties.author = 'Ty Tuff'
doc.save(ROOT / 'cubedynamics-joss-collaborator-draft.docx')
print(f'Created collaborator copy with {len(keys)} resolved references')
